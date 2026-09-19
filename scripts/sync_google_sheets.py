"""Convert a Drive XLSX export to Markdown without saving the intermediate file."""
from datetime import date, datetime, time
from html import escape
from io import BytesIO
from urllib.parse import quote

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def markdown_text(value):
    """Preserve literal cell text without letting it break Markdown tables."""
    text = escape(str(value), quote=False).replace("\\", "\\\\")
    for char in ("|", "*", "_", "`", "[", "]", "#"):
        text = text.replace(char, "\\" + char)
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")


def cell_text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    # Use numeric values rather than rounding metrics to their display precision.
    # Retain percentage units when a cell is formatted as a percentage.
    if isinstance(value, (int, float)) and "%" in cell.number_format:
        return f"{value * 100:.15g}%"
    return str(value)


def workbook_markdown(payload, title):
    formulas = load_workbook(BytesIO(payload), data_only=False)
    values = load_workbook(BytesIO(payload), data_only=True)
    sections = ["# " + markdown_text(title), "",
                "Export: visible worksheets only. Formula results are exported as values; dates use ISO format.",
                "Column letters and original row numbers preserve cell positions; header rows remain in their original rows.",
                "Visual styling, charts/images, and threaded discussions are not represented. Cell notes and merged ranges are listed where available.", ""]
    report = []
    try:
        for sheet in formulas.worksheets:
            if sheet.sheet_state != "visible":
                continue
            cached = values[sheet.title]
            cells, notes, warnings = {}, [], []
            # Sparse cells avoid iterating huge formatted-but-empty rectangular grids.
            # openpyxl's sparse store is covered by tests and constrained to version 3.x.
            for (row, col), cell in sorted(sheet._cells.items()):
                if cell.value is None and cell.comment is None and cell.hyperlink is None:
                    continue
                rendered = cell_text(cached.cell(row, col))
                if cell.data_type == "f" and cached.cell(row, col).value is None:
                    rendered = "[No cached result; formula] " + str(cell.value)
                    warnings.append(f"{cell.coordinate}: formula has no cached value; formula retained.")
                if cell.hyperlink and cell.hyperlink.target:
                    rendered += " (link: " + cell.hyperlink.target + ")"
                cells[(row, col)] = rendered
                if cell.comment:
                    notes.append(f"- {cell.coordinate}: {markdown_text(cell.comment.text)}")
                if cached.cell(row, col).data_type == "e":
                    warnings.append(f"{cell.coordinate}: spreadsheet error value {rendered}.")
            sections += ["## Sheet: " + markdown_text(sheet.title), ""]
            if not cells:
                sections += ["Used range: empty", "", "No populated cells.", ""]
                report.append({"name": sheet.title, "used_range": None, "populated_cells": 0, "warnings": []})
                continue
            min_row, max_row = min(r for r, c in cells), max(r for r, c in cells)
            min_col, max_col = min(c for r, c in cells), max(c for r, c in cells)
            used = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
            sections += [f"Used range: `{used}`", ""]
            merged = sorted(str(item) for item in sheet.merged_cells.ranges)
            if merged:
                sections += ["Merged ranges (value remains at top-left cell): " + ", ".join(merged), ""]
                warnings.append("Merged cell spans are described rather than visually reproduced.")
            active_rows = sorted({r for r, c in cells})
            table = (max_col - min_col + 1 <= 12 and
                     max(len(value) for value in cells.values()) <= 250 and
                     (max_row - min_row + 1) <= len(active_rows) * 3)
            if table:
                columns = list(range(min_col, max_col + 1))
                sections += ["| Row | " + " | ".join(get_column_letter(c) for c in columns) + " |",
                             "| --- | " + " | ".join("---" for c in columns) + " |"]
                for row in range(min_row, max_row + 1):
                    sections.append(f"| {row} | " + " | ".join(markdown_text(cells.get((row, c), "")) for c in columns) + " |")
                sections.append("")
            else:
                sections += ["Structured rows (wide, long-text, or sparse tab). Unlisted cells are blank; no populated columns are omitted.", ""]
                for row in active_rows:
                    sections += [f"### Row {row}", ""]
                    for (r, col), value in cells.items():
                        if r == row:
                            sections.append(f"- **{get_column_letter(col)}{row}**: {markdown_text(value)}")
                    sections.append("")
            if notes:
                sections += ["### Cell notes", ""] + notes + [""]
            if warnings:
                sections += ["### Export warnings", ""] + ["- " + markdown_text(w) for w in warnings] + [""]
            report.append({"name": sheet.title, "used_range": used, "populated_cells": len(cells),
                           "representation": "table" if table else "structured_rows", "warnings": warnings})
    finally:
        formulas.close()
        values.close()
    return "\n".join(sections).rstrip() + "\n", report


def fetch_sheet_markdown(session, source, check_response, error_class):
    base = "https://www.googleapis.com/drive/v3/files/" + quote(source["drive_id"], safe="")
    response = session.get(base, params={"fields": "id,name,mimeType,trashed,webViewLink", "supportsAllDrives": "true"}, timeout=60)
    check_response(response, "Sheet metadata fetch")
    metadata = response.json()
    if metadata.get("mimeType") != "application/vnd.google-apps.spreadsheet" or metadata.get("trashed"):
        raise error_class("Registered source is not a live Google Sheet.")
    response = session.get(base + "/export", params={"mimeType": XLSX_MIME}, timeout=60)
    check_response(response, "Sheet workbook export")
    return workbook_markdown(response.content, source["title"])
