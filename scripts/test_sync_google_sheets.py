"""Offline workbook export tests. No API access or real source content."""
from io import BytesIO
import unittest
import re
from unittest.mock import Mock
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from sync_google_sheets import workbook_markdown, fetch_sheet_markdown, XLSX_MIME
from sync_google_docs import SyncError, check_response


def exported(workbook):
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class SheetTests(unittest.TestCase):
    def test_visible_tabs_used_range_headers_notes_and_escaping(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Metrics"
        sheet.append(["Experiment", "Recall", "Conclusion"])
        sheet.append(["E1", 0.75, "a|b\nnext <step>"])
        sheet['B2'].number_format = '0.0%'
        sheet['C2'].comment = Comment('Analyst judgment', 'Tester')
        sheet['XFD10000'].font = Font(bold=True)
        book.create_sheet('Private').sheet_state = 'hidden'
        book.create_sheet('Empty')
        text, report = workbook_markdown(exported(book), 'Dashboard')
        self.assertEqual([r['name'] for r in report], ['Metrics', 'Empty'])
        self.assertEqual(report[0]['used_range'], 'A1:C2')
        self.assertIn('| 1 | Experiment | Recall | Conclusion |', text)
        self.assertIn('75%', text)
        self.assertIn('a\\|b<br>next &lt;step&gt;', text)
        self.assertIn('Analyst judgment', text)
        self.assertNotIn('Private', text)
        self.assertEqual(report[1]['used_range'], None)

    def test_wide_sparse_rows_keep_last_column_and_zero_false(self):
        book = Workbook()
        sheet = book.active
        sheet['B3'] = 0
        sheet['N3'] = False
        sheet['N1000'] = 'Last column content'
        text, report = workbook_markdown(exported(book), 'Wide')
        self.assertEqual(report[0]['used_range'], 'B3:N1000')
        self.assertEqual(report[0]['representation'], 'structured_rows')
        self.assertIn('**B3**: 0', text)
        self.assertIn('**N3**: FALSE', text)
        self.assertIn('**N1000**: Last column content', text)
        self.assertNotIn('### Row 999', text)

    def test_formulas_cached_values_and_missing_cache_warning(self):
        book = Workbook()
        book.active['A1'] = '=1+1'
        book.active['B1'] = '=2+2'
        payload = exported(book)
        out = BytesIO()
        with ZipFile(BytesIO(payload)) as source, ZipFile(out, 'w', ZIP_DEFLATED) as target:
            for name in source.namelist():
                content = source.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    content = re.sub(rb'(<f>1\+1</f>)<v(?:></v>|\s*/>)', rb'\1<v>2</v>', content)
                target.writestr(name, content)
        text, report = workbook_markdown(out.getvalue(), 'Formulas')
        self.assertIn('| 1 | 2 |', text)
        self.assertNotIn('=1+1', text)
        self.assertIn('=2+2', text)
        self.assertEqual(len(report[0]['warnings']), 1)

    def test_merge_and_long_notes_are_not_lost(self):
        book = Workbook()
        sheet = book.active
        sheet.merge_cells('A1:C1')
        sheet['A1'] = 'Study overview'
        sheet['C4'] = 'Long hypothesis ' * 30
        text, report = workbook_markdown(exported(book), 'Notes')
        self.assertIn('A1:C1', text)
        self.assertIn('**C4**: ' + 'Long hypothesis ' * 30, text)
        self.assertEqual(report[0]['representation'], 'structured_rows')

    def test_drive_export_uses_existing_session(self):
        book = Workbook()
        book.active['A1'] = 'Data'
        session = Mock()
        session.get.side_effect = [Mock(status_code=200, json=lambda: {'mimeType': 'application/vnd.google-apps.spreadsheet'}),
                                   Mock(status_code=200, content=exported(book))]
        text, report = fetch_sheet_markdown(session, {'drive_id': 'selected', 'title': 'Test'}, check_response, SyncError)
        self.assertEqual(session.get.call_count, 2)
        self.assertEqual(session.get.call_args.kwargs['params'], {'mimeType': XLSX_MIME})
        self.assertIn('Data', text)


if __name__ == '__main__':
    unittest.main()
